from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
bullet_slide_layout = prs.slide_layouts[1]

# -------------------------------------------------------------
# SLIDE 1: Module 5 — Vision-Language Intelligence
# -------------------------------------------------------------
slide1 = prs.slides.add_slide(bullet_slide_layout)
title1 = slide1.shapes.title
title1.text = "Module 5 — Vision-Language Intelligence (VQA)"

tf1 = slide1.shapes.placeholders[1].text_frame
tf1.text = "Fulfills Phase 2 goal: Deep Contextual Understanding"
tf1.add_paragraph().text = "Integrates the BLIP Vision-Language Transformer."
tf1.add_paragraph().text = "Upgrades basic object detection to semantic scene understanding."
tf1.add_paragraph().text = "Enables Voice-to-Text interactive Q&A for the user."
tf1.add_paragraph().text = "Processes raw camera pixels alongside voice prompts."
tf1.add_paragraph().text = "Generates context-aware, natural English sentences."

notes_slide1 = slide1.notes_slide
notes_slide1.notes_text_frame.text = "SPEAKER NOTES:\nAs seen in our Phase 1 planned goals, we successfully integrated a Vision-Language Model. While YOLO just detects objects, BLIP acts as the brain. It allows the blind user to pause and ask complex questions using their voice, and the AI generates a context-aware sentence in return."


# -------------------------------------------------------------
# SLIDE 2: Enhanced Multimodal Fusion Engine
# -------------------------------------------------------------
slide2 = prs.slides.add_slide(bullet_slide_layout)
title2 = slide2.shapes.title
title2.text = "Enhanced Multimodal Fusion Engine"

tf2 = slide2.shapes.placeholders[1].text_frame
tf2.text = "Acts as the system's central intelligent decision hub."
tf2.add_paragraph().text = "Fuses Object (YOLO), Depth, Text (OCR), and VQA data."
tf2.add_paragraph().text = "Translates 2D bounding boxes into 3D physical proximity."
tf2.add_paragraph().text = "Assigns strict hazard levels: SAFE, CLOSE, or CRITICAL."
tf2.add_paragraph().text = "Filters duplicate detections to prevent audio overload."

notes_slide2 = slide2.notes_slide
notes_slide2.notes_text_frame.text = "SPEAKER NOTES:\nThis is the traffic cop of our upgraded system. It takes bounding boxes from YOLO, matches them with pixel intensity from the Depth model, and decides the hazard level. It prioritizes Critical hazards and suppresses duplicate alerts so the user isn't overwhelmed by overlapping audio."


# -------------------------------------------------------------
# SLIDE 3: Module 6 — Neural TTS & Edge Deployment
# -------------------------------------------------------------
slide3 = prs.slides.add_slide(bullet_slide_layout)
title3 = slide3.shapes.title
title3.text = "Module 6 — Neural TTS & Edge Deployment"

tf3 = slide3.shapes.placeholders[1].text_frame
tf3.text = "Fulfills Phase 2 goal: Edge / Mobile Deployment"
tf3.add_paragraph().text = "Converts fused JSON hazard warnings into real-time audio."
tf3.add_paragraph().text = "Deployed via local cross-network HTTP API bridge."
tf3.add_paragraph().text = "Transforms a standard smartphone into a wearable sensor."
tf3.add_paragraph().text = "Employs state-based audio cooldown timers."
tf3.add_paragraph().text = "Maintains strict low-latency edge device performance."

notes_slide3 = slide3.notes_slide
notes_slide3.notes_text_frame.text = "SPEAKER NOTES:\nTo fulfill our final Phase 2 goal of mobile deployment, we built a local Wi-Fi API bridge. Because blind users rely entirely on audio, we optimized this Neural TTS to run with extremely low latency, turning the user's smartphone into a completely hands-free wearable sensor."


# -------------------------------------------------------------
# SLIDE 4: Final System & Live Demonstration
# -------------------------------------------------------------
slide4 = prs.slides.add_slide(bullet_slide_layout)
title4 = slide4.shapes.title
title4.text = "Final System Output & Live Demonstration"

tf4 = slide4.shapes.placeholders[1].text_frame
tf4.text = "All Phase 2 objectives have been successfully implemented."
tf4.add_paragraph().text = "Seamlessly bridges local inference with mobile UI."
tf4.add_paragraph().text = "Provides true hands-free, interactive blind navigation."
tf4.add_paragraph().text = ""
tf4.add_paragraph().text = "Switching to Live Mobile Camera Demonstration..."

notes_slide4 = slide4.notes_slide
notes_slide4.notes_text_frame.text = "SPEAKER NOTES:\nWe have successfully transformed our Phase 1 prototype into a fully unified, real-time edge application. We will now demonstrate the system live. I am placing my smartphone in my pocket to act as the user's camera, and we will show you all 6 models working together in real-time."

prs.save("Avinash_Final_Phase2_Slides.pptx")
