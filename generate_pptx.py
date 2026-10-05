from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()

# Use layout 1 (Title and Content)
bullet_slide_layout = prs.slide_layouts[1]

# -------------------------------------------------------------
# SLIDE 1: Model 5 (BLIP)
# -------------------------------------------------------------
slide1 = prs.slides.add_slide(bullet_slide_layout)
title1 = slide1.shapes.title
title1.text = "Model 5: Vision-Language Intelligence (BLIP)"

content1 = slide1.shapes.placeholders[1]
tf1 = content1.text_frame
tf1.text = "Utilizes BLIP Vision-Language Transformer architecture."

p = tf1.add_paragraph()
p.text = "Upgrades basic detection to full contextual scene understanding."
p = tf1.add_paragraph()
p.text = "Processes raw image pixels alongside user voice prompts."
p = tf1.add_paragraph()
p.text = "Enables interactive Voice-to-Text Q&A for the blind user."
p = tf1.add_paragraph()
p.text = "Generates complete, natural English sentences natively."
p = tf1.add_paragraph()
p.text = "Resolves complex spatial queries (e.g., \"Is this seat empty?\")."

notes_slide1 = slide1.notes_slide
text_frame1 = notes_slide1.notes_text_frame
text_frame1.text = "SPEAKER NOTES:\nWhile YOLO just detects objects, our BLIP Vision-Language model acts as the system's brain. It allows the blind user to pause and ask complex questions about their environment. We integrated a voice-activated Q&A system where the user speaks, and the AI generates a context-aware sentence."

# -------------------------------------------------------------
# SLIDE 2: Multimodal Fusion Engine
# -------------------------------------------------------------
slide2 = prs.slides.add_slide(bullet_slide_layout)
title2 = slide2.shapes.title
title2.text = "The Core: Multimodal Fusion Engine"

content2 = slide2.shapes.placeholders[1]
tf2 = content2.text_frame
tf2.text = "Acts as the system's central intelligent hub."

p = tf2.add_paragraph()
p.text = "Synchronizes Object (YOLO), Depth, and Text (OCR) data."
p = tf2.add_paragraph()
p.text = "Maps 2D bounding boxes directly to 3D depth pixels."
p = tf2.add_paragraph()
p.text = "Calculates exact proximity to assign hazard priority levels."
p = tf2.add_paragraph()
p.text = "Tags all obstacles strictly as SAFE, CLOSE, or CRITICAL."
p = tf2.add_paragraph()
p.text = "Suppresses duplicate warnings to prevent audio spam."

notes_slide2 = slide2.notes_slide
text_frame2 = notes_slide2.notes_text_frame
text_frame2.text = "SPEAKER NOTES:\nThe Multimodal Fusion Engine is the traffic cop of our system. It takes the bounding boxes from YOLO, matches them with the pixel intensity from the Depth model, and decides how dangerous the obstacle is. It prioritizes Critical hazards and suppresses duplicate alerts so the user isn't overwhelmed by overlapping audio."

# -------------------------------------------------------------
# SLIDE 3: Neural TTS (Model 6)
# -------------------------------------------------------------
slide3 = prs.slides.add_slide(bullet_slide_layout)
title3 = slide3.shapes.title
title3.text = "Model 6: Neural TTS & Edge Delivery"

content3 = slide3.shapes.placeholders[1]
tf3 = content3.text_frame
tf3.text = "Converts fused JSON text data into real-time audio."

p = tf3.add_paragraph()
p.text = "Utilizes high-speed Neural Text-to-Speech synthesis."
p = tf3.add_paragraph()
p.text = "Employs intelligent, state-based audio-cooldown timers."
p = tf3.add_paragraph()
p.text = "Operates seamlessly via local Wi-Fi API bridging."
p = tf3.add_paragraph()
p.text = "Maintains strict low-latency edge device performance."
p = tf3.add_paragraph()
p.text = "Delivers completely hands-free, wearable blind navigation."

notes_slide3 = slide3.notes_slide
text_frame3 = notes_slide3.notes_text_frame
text_frame3.text = "SPEAKER NOTES:\nOur final stage is Model 6, the Neural Text-to-Speech engine. Because blind users rely entirely on audio, we optimized this TTS to run with extremely low latency over a local Wi-Fi bridge to their smartphone. It uses smart cooldown timers so it speaks smoothly, providing true hands-free navigation."

prs.save("Avinash_Part4_Slides.pptx")
